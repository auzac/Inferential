"""Reading CSV and Excel files into the raw data layer.

Nothing here guesses silently: delimiter/encoding/decimal detection only *proposes* options
that the user confirms in the import wizard, timestamps are parsed with an explicit
:class:`TimestampSpec`, and every cell that is not a number is counted and reported rather
than quietly becoming NaN.
"""

from __future__ import annotations

import codecs
import csv
import itertools
import re
from dataclasses import dataclass, replace
from pathlib import Path

import numpy as np
import pandas as pd

from inferential_studio.core.dataset import (
    ColumnQuality,
    DataImportError,
    Gap,
    ImportReport,
    Issue,
    ProjectData,
    Severity,
    SourceData,
)
from inferential_studio.core.persistence import resolve_source_path, sha256_file
from inferential_studio.core.schemas import (
    DataSource,
    DuplicatePolicy,
    Project,
    SourceKind,
    TimestampKind,
)
from inferential_studio.core.tasks import CancelToken, ProgressCallback, check, report
from inferential_studio.core.timestamps import ParsedTimestamps, TimestampError, parse_timestamps

CSV_SUFFIXES = {".csv", ".tsv", ".txt"}
EXCEL_SUFFIXES = {".xlsx", ".xlsm"}
LARGE_FILE_BYTES = 1024**3
CSV_CHUNK_ROWS = 200_000
MAX_GAPS_LISTED = 20
# Text that means "no value" rather than an informative status token.
MISSING_TOKENS = frozenset(
    {"nan", "NaN", "NAN", "NA", "N/A", "n/a", "#N/A", "null", "NULL", "None", "-"}
)


_COMMA_NUMBER = re.compile(r"^[-+]?\d+,\d+$")
_DOT_NUMBER = re.compile(r"^[-+]?\d+\.\d+$")


class SourceChangedError(DataImportError):
    """The file's checksum differs from the one recorded at import."""


@dataclass(frozen=True)
class ReadOptions:
    delimiter: str | None = None  # CSV only
    encoding: str = "utf-8"  # CSV only
    decimal: str = "."  # CSV only (Excel stores real numbers)
    header_row: int = 0  # 0-based line/row holding the column names
    sheet: str | None = None  # Excel only

    @classmethod
    def from_source(cls, source: DataSource) -> ReadOptions:
        return cls(
            delimiter=source.delimiter,
            encoding=source.encoding,
            decimal=source.decimal,
            header_row=source.header_row,
            sheet=source.sheet,
        )


# --------------------------------------------------------------------------- file inspection


def file_kind(path: str | Path) -> str:
    suffix = Path(path).suffix.lower()
    if suffix in CSV_SUFFIXES:
        return "csv"
    if suffix in EXCEL_SUFFIXES:
        return "excel"
    if suffix == ".xls":
        raise DataImportError(
            "Legacy .xls files are not supported. Open the file in Excel and save it as .xlsx."
        )
    raise DataImportError(
        f"Unsupported file type '{suffix}'. Use CSV (.csv, .tsv, .txt) or Excel (.xlsx)."
    )


def sniff_csv(path: str | Path, sample_bytes: int = 256 * 1024) -> ReadOptions:
    """Propose encoding, delimiter and decimal separator for a delimited text file."""
    with open(path, "rb") as handle:
        raw = handle.read(sample_bytes)
    encoding = "utf-8"
    if raw.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        encoding = "utf-16"
        text = raw.decode("utf-16", errors="ignore")
        return _propose(path, text, encoding)
    try:
        codecs.getincrementaldecoder("utf-8-sig")().decode(raw, final=False)
        text = raw.decode("utf-8-sig", errors="ignore")
        if raw.startswith(codecs.BOM_UTF8):
            encoding = "utf-8-sig"
    except UnicodeDecodeError:
        encoding = "cp1252"  # the usual Windows "ANSI" encoding of historian exports
        text = raw.decode(encoding, errors="replace")

    return _propose(path, text, encoding)


def _propose(path: str | Path, text: str, encoding: str) -> ReadOptions:
    lines = [ln for ln in text.splitlines()[:50] if ln.strip()]
    sample = "\n".join(lines)
    delimiter = ","
    if Path(path).suffix.lower() == ".tsv":
        delimiter = "\t"
    elif sample:
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
        except csv.Error:
            counts = {d: lines[0].count(d) for d in ",;\t|"}
            delimiter = max(counts, key=counts.get) if max(counts.values()) else ","

    decimal = "."
    if delimiter != ",":
        # Look only at fields that are entirely a number, so dates like 01.02.2026 don't count.
        fields = [f.strip() for ln in lines[1:] for f in ln.split(delimiter)]
        comma = sum(bool(_COMMA_NUMBER.match(f)) for f in fields)
        dot = sum(bool(_DOT_NUMBER.match(f)) for f in fields)
        if comma > dot:
            decimal = ","
    return ReadOptions(delimiter=delimiter, encoding=encoding, decimal=decimal)


def list_sheets(path: str | Path) -> list[str]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(path, read_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl raises many types for bad files
        raise DataImportError(
            f"'{Path(path).name}' cannot be opened as an Excel workbook: {exc}"
        ) from exc
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def excel_date_system(path: str | Path) -> int:
    """1900 or 1904 — needed to interpret numeric (serial) dates stored in a workbook."""
    from openpyxl import load_workbook
    from openpyxl.utils.datetime import CALENDAR_MAC_1904

    wb = load_workbook(path, read_only=True)
    try:
        return 1904 if wb.epoch == CALENDAR_MAC_1904 else 1900
    finally:
        wb.close()


def _csv_kwargs(opts: ReadOptions) -> dict:
    return {
        "sep": opts.delimiter or ",",
        "encoding": opts.encoding,
        "skip_blank_lines": False,  # keep line numbers aligned with the file
        "engine": "c",
    }


def _encoding_error(path: str | Path, opts: ReadOptions) -> DataImportError:
    return DataImportError(
        f"'{Path(path).name}' is not valid {opts.encoding} text. Choose another encoding (e.g. cp1252)."
    )


def read_header(path: str | Path, opts: ReadOptions) -> list[str]:
    """Column names exactly as written, rejecting blank and duplicate names."""
    try:
        if file_kind(path) == "csv":
            # csv module, not pandas: preamble lines above the header may have any number of fields.
            with open(path, encoding=opts.encoding, newline="") as handle:
                rows = list(
                    itertools.islice(
                        csv.reader(handle, delimiter=opts.delimiter or ","), opts.header_row + 1
                    )
                )
            header = rows[opts.header_row] if len(rows) > opts.header_row else None
        else:
            head = pd.read_excel(
                path,
                sheet_name=opts.sheet or 0,
                header=None,
                nrows=opts.header_row + 1,
                dtype=object,
            )
            header = list(head.iloc[opts.header_row]) if len(head) > opts.header_row else None
    except UnicodeDecodeError as exc:
        raise _encoding_error(path, opts) from exc
    except (csv.Error, ValueError) as exc:
        raise DataImportError(f"Cannot read the header of '{Path(path).name}': {exc}") from exc
    if header is None:
        raise DataImportError(f"'{Path(path).name}' has fewer than {opts.header_row + 1} lines.")
    names = [
        "" if v is None or (isinstance(v, float) and np.isnan(v)) else str(v).strip()
        for v in header
    ]
    blanks = [i + 1 for i, n in enumerate(names) if not n]
    if blanks:
        raise DataImportError(
            f"Column(s) {', '.join(map(str, blanks))} have no name in the header row. "
            "Name every column (or choose the correct header row)."
        )
    seen, dupes = set(), []
    for n in names:
        if n in seen and n not in dupes:
            dupes.append(n)
        seen.add(n)
    if dupes:
        raise DataImportError(
            f"Duplicate column name(s): {', '.join(dupes)}. Rename them in the file so each column is unique."
        )
    return names


def read_preview(path: str | Path, opts: ReadOptions, nrows: int = 200) -> pd.DataFrame:
    """First rows with original column names, for the wizard's preview table."""
    names = read_header(path, opts)
    if file_kind(path) == "csv":
        df = pd.read_csv(
            path,
            header=None,
            skiprows=opts.header_row + 1,
            nrows=nrows,
            dtype=str,
            names=names,
            keep_default_na=False,
            **_csv_kwargs(opts),
        )
    else:
        df = pd.read_excel(
            path,
            sheet_name=opts.sheet or 0,
            header=None,
            skiprows=opts.header_row + 1,
            nrows=nrows,
            names=names,
        )
    return df


def read_column(path: str | Path, opts: ReadOptions, column: str) -> pd.Series:
    """Every value of one column, raw (text for CSV; cell values for Excel).

    Used by the wizard so timestamp-format detection can look for a value that settles
    day/month order anywhere in the file, not just in the preview rows.
    """
    names = read_header(path, opts)
    if column not in names:
        raise DataImportError(f"Column '{column}' not found.")
    if file_kind(path) == "csv":
        try:
            return pd.read_csv(
                path,
                header=None,
                skiprows=opts.header_row + 1,
                names=names,
                usecols=[column],
                dtype=str,
                keep_default_na=False,
                **_csv_kwargs(opts),
            )[column]
        except UnicodeDecodeError as exc:
            raise _encoding_error(path, opts) from exc
    return pd.read_excel(
        path, sheet_name=opts.sheet or 0, header=opts.header_row, usecols=[column]
    )[column]


# --------------------------------------------------------------------------- full import


def import_source(
    source: DataSource,
    project_path: str | Path | None = None,
    *,
    value_columns: list[str] | None = None,
    verify_checksum: bool = True,
    progress: ProgressCallback | None = None,
    cancel: CancelToken | None = None,
) -> SourceData:
    """Read one source file into the raw layer and build its validation report.

    ``value_columns`` defaults to every column other than the timestamp/id columns.
    """
    path = resolve_source_path(source, project_path)
    if not path.is_file():
        raise DataImportError(f"File '{source.path}' was not found.")
    report(progress, 0.0, f"Checking {path.name}")
    digest = sha256_file(path)
    if verify_checksum and source.sha256 and digest != source.sha256:
        raise SourceChangedError(
            f"'{path.name}' has changed since it was imported (checksum mismatch). "
            "Re-import it to accept the new contents."
        )
    check(cancel)

    opts = ReadOptions.from_source(source)
    names = read_header(path, opts)
    special = list(source.timestamp_columns)
    missing = [c for c in special + (value_columns or []) if c not in names]
    if missing:
        raise DataImportError(
            f"'{path.name}' has no column(s) {', '.join(repr(m) for m in missing)}. "
            f"Available: {', '.join(names)}."
        )
    if value_columns is None:
        value_columns = [n for n in names if n not in special]

    def read_progress(f: float, m: str) -> None:
        report(progress, 0.05 + 0.75 * f, m)

    if file_kind(path) == "csv":
        raw_special, values, text_mask, quality = _read_csv(
            path, opts, names, special, value_columns, read_progress, cancel
        )
    else:
        raw_special, values, text_mask, quality = _read_excel(
            path, opts, special, value_columns, cancel
        )
    check(cancel)
    report(progress, 0.85, "Parsing timestamps")

    try:
        parsed = parse_timestamps(raw_special[source.timestamp.column], source.timestamp)
        result_parsed = (
            parse_timestamps(raw_special[source.result_timestamp.column], source.result_timestamp)
            if source.result_timestamp
            else None
        )
    except TimestampError as exc:
        raise DataImportError(str(exc)) from exc

    n = len(values)
    file_line = np.arange(n) + opts.header_row + 2
    sample_id = (
        raw_special[source.sample_id_column].astype("string").str.strip()
        if source.sample_id_column
        else None
    )
    report(progress, 0.92, "Checking data quality")
    rep = build_report(
        source, file_line, raw_special[source.timestamp.column], parsed, result_parsed, quality
    )
    data = SourceData(
        source=replace(source, sha256=digest),
        sha256=digest,
        file_line=file_line,
        raw_time_text=raw_special[source.timestamp.column].astype("string"),
        time=parsed.utc,
        values=values,
        text_mask=text_mask,
        report=rep,
        result_time=result_parsed.utc if result_parsed is not None else None,
        sample_id=sample_id,
    )
    report(progress, 1.0, "Done")
    return data


def _to_numeric(raw: pd.Series, decimal: str) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    """Return (values, blank_mask, text_mask) for one column of raw cells."""
    if pd.api.types.is_numeric_dtype(raw) and not pd.api.types.is_bool_dtype(raw):
        values = np.array(raw, dtype="float64")
        blank = np.isnan(values)
        return values, blank, np.zeros(len(raw), dtype=bool)
    text = raw.astype("string").str.strip()
    blank = (text.isna() | (text == "") | text.isin(MISSING_TOKENS)).to_numpy(dtype=bool)
    candidate = text.str.replace(",", ".", regex=False) if decimal == "," else text
    values = np.array(pd.to_numeric(candidate, errors="coerce"), dtype="float64")
    values[blank] = np.nan
    textual = ~blank & np.isnan(values)
    return values, blank, textual


def _accumulate_quality(
    acc: dict[str, ColumnQuality], name: str, raw: pd.Series, values, blank, textual
) -> None:
    q = acc.setdefault(name, ColumnQuality(name, 0, 0, 0, {}))
    q.n_blank += int(blank.sum())
    q.n_text += int(textual.sum())
    q.n_numeric += int((~blank & ~textual).sum())
    if textual.any():
        tokens = raw.astype("string").str.strip()[textual].value_counts().head(5)
        for token, count in tokens.items():
            q.text_examples[str(token)] = q.text_examples.get(str(token), 0) + int(count)


def _finalize_quality(acc: dict[str, ColumnQuality], columns: list[str]) -> list[ColumnQuality]:
    out = []
    for c in columns:
        q = acc[c]
        q.text_examples = dict(sorted(q.text_examples.items(), key=lambda kv: -kv[1])[:5])
        out.append(q)
    return out


def _read_csv(path, opts, names, special, value_columns, progress, cancel):
    usecols = special + value_columns
    file_size = max(Path(path).stat().st_size, 1)
    specials: dict[str, list[pd.Series]] = {c: [] for c in special}
    value_parts: list[np.ndarray] = []
    mask_parts: list[np.ndarray] = []
    acc: dict[str, ColumnQuality] = {}
    try:
        with open(path, "rb") as handle:
            reader = pd.read_csv(
                handle,
                header=None,
                skiprows=opts.header_row + 1,
                names=names,
                usecols=usecols,
                # Timestamp/id columns stay text; value columns use the fast C float parser and
                # fall back to cell-by-cell checking only when a chunk contains text tokens.
                dtype={c: str for c in special},
                keep_default_na=False,
                na_values=sorted(MISSING_TOKENS | {""}),
                decimal=opts.decimal,
                chunksize=CSV_CHUNK_ROWS,
                **_csv_kwargs(opts),
            )
            for chunk in reader:
                check(cancel)
                for c in special:
                    specials[c].append(chunk[c])
                vals = np.empty((len(chunk), len(value_columns)))
                mask = np.zeros((len(chunk), len(value_columns)), dtype=bool)
                for j, c in enumerate(value_columns):
                    v, blank, textual = _to_numeric(chunk[c], opts.decimal)
                    vals[:, j], mask[:, j] = v, textual
                    _accumulate_quality(acc, c, chunk[c], v, blank, textual)
                value_parts.append(vals)
                mask_parts.append(mask)
                progress(min(handle.tell() / file_size, 1.0), f"Reading {Path(path).name}")
    except UnicodeDecodeError as exc:
        raise _encoding_error(path, opts) from exc
    except pd.errors.ParserError as exc:
        raise DataImportError(
            f"'{Path(path).name}' could not be parsed with the chosen delimiter: {exc}"
        ) from exc

    if not value_parts:
        raise DataImportError(f"'{Path(path).name}' has no data rows below the header.")
    raw_special = pd.DataFrame(
        {c: pd.concat(parts, ignore_index=True) for c, parts in specials.items()}
    )
    for c in value_columns:
        acc.setdefault(c, ColumnQuality(c, 0, 0, 0, {}))
    values = pd.DataFrame(np.vstack(value_parts), columns=value_columns)
    text_mask = pd.DataFrame(np.vstack(mask_parts), columns=value_columns)
    return raw_special, values, text_mask, _finalize_quality(acc, value_columns)


def _read_excel(path, opts, special, value_columns, cancel):
    try:
        df = pd.read_excel(
            path,
            sheet_name=opts.sheet or 0,
            header=opts.header_row,
            usecols=special + value_columns,
        )
    except ValueError as exc:
        raise DataImportError(
            f"Cannot read sheet '{opts.sheet}' of '{Path(path).name}': {exc}"
        ) from exc
    check(cancel)
    if df.empty:
        raise DataImportError(f"'{Path(path).name}' has no data rows below the header.")
    acc: dict[str, ColumnQuality] = {}
    vals = np.empty((len(df), len(value_columns)))
    mask = np.zeros((len(df), len(value_columns)), dtype=bool)
    for j, c in enumerate(value_columns):
        v, blank, textual = _to_numeric(df[c], ".")
        vals[:, j], mask[:, j] = v, textual
        _accumulate_quality(acc, c, df[c], v, blank, textual)
    raw_special = df[special].reset_index(drop=True)
    return (
        raw_special,
        pd.DataFrame(vals, columns=value_columns),
        pd.DataFrame(mask, columns=value_columns),
        _finalize_quality(acc, value_columns),
    )


# --------------------------------------------------------------------------- validation report


def build_report(
    source: DataSource,
    file_line: np.ndarray,
    raw_time: pd.Series,
    parsed: ParsedTimestamps,
    result_parsed: ParsedTimestamps | None,
    columns: list[ColumnQuality],
) -> ImportReport:
    times = parsed.utc
    valid = times.notna().to_numpy()
    valid_times = times[valid]
    invalid_rows = np.flatnonzero(
        (parsed.unparseable | parsed.dst_ambiguous | parsed.dst_nonexistent).to_numpy()
    )
    examples = [(int(file_line[i]), str(raw_time.iloc[i])) for i in invalid_rows[:5]]

    tvals = valid_times.dt.tz_convert(None).to_numpy(dtype="datetime64[ns]")  # naive UTC
    n_out_of_order = int((np.diff(tvals) < np.timedelta64(0, "ns")).sum()) if len(tvals) > 1 else 0
    n_dupes = int(pd.DatetimeIndex(valid_times).duplicated().sum())

    unique_sorted = np.unique(tvals)
    median = gap_threshold = None
    irregular = 0.0
    gaps: list[Gap] = []
    n_gaps = 0
    if len(unique_sorted) > 2:
        deltas = np.diff(unique_sorted)
        median_ns = np.median(deltas.astype("int64"))
        median = pd.Timedelta(int(median_ns), unit="ns")
        irregular = float(np.mean(np.abs(deltas.astype("int64") - median_ns) > 0.1 * median_ns))
        gap_threshold = max(5 * median, pd.Timedelta(minutes=1))
        long = np.flatnonzero(deltas > gap_threshold.to_timedelta64())
        n_gaps = int(long.size)
        longest = long[np.argsort(deltas[long])[::-1][:MAX_GAPS_LISTED]]
        gaps = [
            Gap(
                pd.Timestamp(unique_sorted[i]).tz_localize("UTC"),
                pd.Timestamp(unique_sorted[i + 1]).tz_localize("UTC"),
            )
            for i in longest
        ]

    rep = ImportReport(
        source_id=source.source_id,
        kind=source.kind,
        n_rows=len(times),
        n_valid_times=int(valid.sum()),
        n_blank_times=int(parsed.blank.sum()),
        n_unparseable_times=int(parsed.unparseable.sum()),
        n_dst_ambiguous=int(parsed.dst_ambiguous.sum()),
        n_dst_nonexistent=int(parsed.dst_nonexistent.sum()),
        invalid_time_examples=examples,
        n_duplicate_times=n_dupes,
        n_out_of_order=n_out_of_order,
        time_start=valid_times.min() if valid.any() else None,
        time_end=valid_times.max() if valid.any() else None,
        median_interval=median,
        irregular_fraction=irregular,
        gaps=gaps,
        n_gaps=n_gaps,
        gap_threshold=gap_threshold,
        columns=columns,
    )
    rep.issues = _issues(source, rep, result_parsed, times)
    return rep


def _issues(
    source: DataSource, rep: ImportReport, result: ParsedTimestamps | None, sample_times: pd.Series
) -> list[Issue]:
    out: list[Issue] = []
    is_lab = source.kind is SourceKind.LAB

    def add(sev: Severity, code: str, msg: str) -> None:
        out.append(Issue(sev, code, msg))

    if rep.n_valid_times == 0:
        add(
            Severity.ERROR,
            "no_valid_times",
            "No row has a valid timestamp; check the timestamp column and format.",
        )
    bad = rep.n_unparseable_times
    if bad:
        ex = "; ".join(f"line {line}: '{text}'" for line, text in rep.invalid_time_examples[:3])
        add(
            Severity.WARNING,
            "unparseable_times",
            f"{bad} timestamp(s) do not match the chosen format and are excluded ({ex}).",
        )
    if rep.n_dst_ambiguous:
        add(
            Severity.WARNING,
            "dst_ambiguous",
            f"{rep.n_dst_ambiguous} local time(s) fall in the repeated hour when clocks go back and cannot be "
            "placed unambiguously; they are excluded.",
        )
    if rep.n_dst_nonexistent:
        add(
            Severity.WARNING,
            "dst_nonexistent",
            f"{rep.n_dst_nonexistent} local time(s) fall in the skipped hour when clocks go forward (they do not "
            "exist in the source timezone); they are excluded. Check the source timezone.",
        )
    if rep.n_blank_times:
        add(
            Severity.INFO,
            "blank_times",
            f"{rep.n_blank_times} row(s) have no timestamp and are excluded.",
        )
    if rep.n_duplicate_times:
        policy = source.effective_duplicate_policy
        if policy is DuplicatePolicy.ERROR:
            add(
                Severity.ERROR,
                "duplicate_times",
                f"{rep.n_duplicate_times} row(s) repeat an earlier timestamp. Choose to keep the first row or "
                "average rows with the same time.",
            )
        else:
            label = {
                DuplicatePolicy.KEEP_FIRST: "only the first row is kept",
                DuplicatePolicy.MEAN: "values are averaged",
                DuplicatePolicy.KEEP_ALL: "rows are kept as separate samples",
            }[policy]
            add(
                Severity.INFO,
                "duplicate_times",
                f"{rep.n_duplicate_times} row(s) repeat an earlier timestamp; {label}.",
            )
    if rep.n_out_of_order:
        add(
            Severity.INFO,
            "out_of_order",
            f"{rep.n_out_of_order} row(s) are out of chronological order; data is sorted by time.",
        )
    if not is_lab and rep.irregular_fraction > 0.05 and rep.median_interval is not None:
        add(
            Severity.WARNING,
            "irregular_sampling",
            f"{rep.irregular_fraction:.0%} of sampling intervals differ by more than 10% from the typical "
            f"interval ({format_duration(rep.median_interval)}).",
        )
    if rep.n_gaps and rep.gaps:
        longest = rep.gaps[0]
        sev = Severity.INFO if is_lab else Severity.WARNING
        add(
            sev,
            "gaps",
            f"{rep.n_gaps} gap(s) longer than {format_duration(rep.gap_threshold)}; the longest is "
            f"{format_duration(longest.duration)} starting {longest.start:%Y-%m-%d %H:%M} UTC.",
        )
    for col in rep.columns:
        if col.n_numeric == 0:
            add(
                Severity.WARNING, "column_empty", f"Column '{col.name}' contains no numeric values."
            )
        elif col.n_text:
            tokens = ", ".join(f"'{t}' ×{n}" for t, n in list(col.text_examples.items())[:3])
            add(
                Severity.WARNING,
                "text_values",
                f"Column '{col.name}' has {col.n_text} non-numeric cell(s) treated as invalid ({tokens}).",
            )
        if is_lab and col.n_blank + col.n_text:
            add(
                Severity.WARNING,
                "missing_target",
                f"{col.n_blank + col.n_text} sample(s) have no valid '{col.name}' value.",
            )

    if result is not None:
        both = result.utc.notna() & sample_times.notna()
        early = int((result.utc[both] < sample_times[both]).sum())
        if early:
            add(
                Severity.WARNING,
                "result_before_sample",
                f"{early} lab result time(s) are earlier than their sample time; check the two columns are not swapped.",
            )
        if result.n_invalid:
            add(
                Severity.WARNING,
                "result_time_invalid",
                f"{result.n_invalid} result time(s) could not be parsed.",
            )
    return out


def format_duration(td: pd.Timedelta | None) -> str:
    if td is None:
        return "?"
    seconds = td.total_seconds()
    if seconds < 120:
        return f"{seconds:g} s"
    if seconds < 7200:
        return f"{seconds / 60:g} min"
    if seconds < 172800:
        return f"{seconds / 3600:.3g} h"
    return f"{seconds / 86400:.3g} days"


def lab_coverage_issue(process: SourceData, lab: SourceData) -> Issue | None:
    """Warn when lab samples fall outside the period covered by process data."""
    p0, p1 = process.report.time_start, process.report.time_end
    if p0 is None or p1 is None:
        return None
    t = lab.time.dropna()
    outside = int(((t < p0) | (t > p1)).sum())
    if outside:
        return Issue(
            Severity.WARNING,
            "lab_outside_process",
            f"{outside} of {len(t)} lab sample(s) in '{lab.source.path}' fall outside the process data period "
            f"({p0:%Y-%m-%d %H:%M} – {p1:%Y-%m-%d %H:%M} UTC). Check both sources' timezones.",
        )
    return None


# --------------------------------------------------------------------------- project loading


def load_project_data(
    project: Project,
    project_path: str | Path | None,
    *,
    verify_checksum: bool = True,
    progress: ProgressCallback | None = None,
    cancel: CancelToken | None = None,
) -> ProjectData:
    """Import every source of a project, reading only the columns its variables use.

    A source that cannot be loaded (missing or changed file, bad format) is recorded in
    ``ProjectData.failures`` instead of aborting the others.
    """
    data = ProjectData()
    n = max(len(project.sources), 1)
    for i, source in enumerate(project.sources):
        columns = [v.source_column for v in project.variables if v.source_id == source.source_id]

        def sub(f: float, m: str, i: int = i) -> None:
            report(progress, (i + f) / n, m)

        try:
            data.sources[source.source_id] = import_source(
                source,
                project_path,
                value_columns=columns,
                verify_checksum=verify_checksum,
                progress=sub,
                cancel=cancel,
            )
        except DataImportError as exc:  # keep loading the other sources; report this one
            data.failures[source.source_id] = exc
    report(progress, 1.0, "Done")
    return data


__all__ = [
    "DataImportError",
    "ReadOptions",
    "SourceChangedError",
    "TimestampKind",
    "excel_date_system",
    "file_kind",
    "import_source",
    "lab_coverage_issue",
    "list_sheets",
    "load_project_data",
    "read_header",
    "read_preview",
    "sniff_csv",
]
